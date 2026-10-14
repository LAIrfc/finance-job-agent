"""Build a source-linked finance recruitment catalogue from preserved research snapshots."""
import csv,hashlib,html,json,re,sqlite3
from collections import Counter
from datetime import date
from pathlib import Path
from urllib.parse import urlsplit,urljoin
from bs4 import BeautifulSoup
ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'data/research/2026-10-13'
OUT=ROOT/'汇总/2026-10-13'; OUT.mkdir(parents=True,exist_ok=True)
TODAY=date(2026,10,13)
SOURCES={s['id']:s for s in json.loads((ROOT/'data/sources.json').read_text())}
PAGES={s['id']:s for s in json.loads((DATA/'research-pages.json').read_text())}
records=[]
def strip(s):return BeautifulSoup(str(s or ''),'html.parser').get_text(' ',strip=True)
def key(s):return hashlib.sha256(s.encode()).hexdigest()[:20]
def sector(name):
 if re.search('中国银保信|中央结算|人民银行|金融监督|证监会|交易所|中国结算|上海清算',name):return '监管与金融基础设施'
 if name in ('中金公司','国信证券','首创证券招聘体系'):return '券商'
 for pattern,tag in [('保险|人保|人寿|太平|太康|再保','保险'),('证券|证劵','券商'),('银行|农商|农信','银行'),('基金|私募|量化|投资|资本|资产|资管|Trading|AIphaGrep','基金/资管/投资'),('期货','期货'),('信托|金租|租赁|理财|金融|消金','其他金融')]:
  if re.search(pattern,name,re.I):return tag
 return '金融科技及其他'
def level(unit):
 if re.search('分行|分公司|支公司|营业部',unit):return '分支机构'
 if re.search('总行|总部|总公司|集团本部',unit):return '总部'
 return '待核实'
def deadline(value):
 try:
  d=date.fromisoformat(str(value)[:10])
  if d.year<2020 or d.year>=2030:return None
  return d.isoformat()
 except ValueError:return None
def add(**r):
 defaults=dict(record_type='具体岗位',org='',org_unit='',org_level='待核实',track='',title='',locations='',recruitment_cycle='届别待核实',deadline=None,deadline_text='',degree='',requirements='',duties='',source_id='',source_url='',apply_url='',verification='',status='',notes='',evidence_file='')
 defaults.update(r);r=defaults
 r['track']=r['track'] or sector(r['org'])
 if not r['status']:
  if r['deadline'] and date.fromisoformat(r['deadline'])<TODAY:r['status']='已截止'
  elif r['recruitment_cycle']=='2026或更早':r['status']='历史届别'
  elif r['deadline']:r['status']='公告期限内（投递前复核）'
  else:r['status']='截止/在招状态待核实'
 r['id']=r.get('id') or key('|'.join([r['org'],r['org_unit'],r['title'],r['locations'],r['recruitment_cycle']]))
 records.append(r)

def expand(table):
 grid={}
 for y,tr in enumerate(table.find_all('tr')):
  x=0
  for td in tr.find_all(['th','td'],recursive=False):
   while (y,x) in grid:x+=1
   rs=int(td.get('rowspan',1));cs=int(td.get('colspan',1));value=td.get_text(' ',strip=True)
   for dy in range(rs):
    for dx in range(cs):grid[y+dy,x+dx]=value
   x+=cs
 if not grid:return []
 return [[grid.get((y,x),'') for x in range(max(k[1] for k in grid)+1)] for y in range(max(k[0] for k in grid)+1)]

def table(sid,index):
 soup=BeautifulSoup((ROOT/PAGES[sid]['raw_file']).read_bytes(),'html.parser')
 return expand(soup.find_all('table')[index])

def page_record(sid,org,title,unit='',loc='',req='',degree='',dd=None,dt='',kind='具体岗位',apply=None,notes=''):
 source=PAGES[sid]; official=not any(x in urlsplit(source['url']).hostname for x in ['.edu.cn','ncss.cn'])
 add(source_id=sid,org=org,org_unit=unit,org_level=level(unit),title=title,locations=loc,requirements=req,degree=degree,deadline=dd,deadline_text=dt or dd or '',record_type=kind,recruitment_cycle='2027校招',source_url=source['url'],apply_url=apply or source['url'],verification='官方公告已读' if official else '高校/就业平台公告已读，需官方复核',notes=notes,evidence_file=source['raw_file'])

# Public API jobs: exact official identifiers, source-specific campus filters, full pagination.
for path in sorted(DATA.glob('zhiye-*-jobs.json')):
 sid=path.name[len('zhiye-'):-len('-jobs.json')];source=SOURCES[sid]
 list_url=source.get('collection_url') or source['url']
 if sid=='rui':list_url='https://ruiinsurance.zhiye.com/custom/b8e2dafd-df98-b35b-ce27-7622bc75e4e7'
 org='中国人保' if sid.startswith('picc-') else source['name']
 for j in json.loads(path.read_text()):
  title=j['JobAdName'];req=strip(j.get('Require'));d=deadline(j.get('EndTime'))
  year='2027校招' if '2027' in title else '2026或更早' if re.search(r'202[0-6]',title) else '校招（届别待核实）'
  unit=j.get('Org') or ''
  link='https://cicc.zhiye.com/custom/campusxq?hideMenu=1&jobAdId='+j['Id'] if sid=='cicc' else list_url
  add(id=key(urlsplit(source['url']).hostname+'|'+j['Id']),source_id=sid,org=org,org_unit=unit,org_level=level(unit),title=title,locations=' / '.join(j.get('LocNames') or []),recruitment_cycle=year,deadline=d,deadline_text=j.get('EndTime') if d else '未公布（系统占位日期不当作真实期限）',requirements=req,duties=strip(j.get('Duty')),degree=str(j.get('Degree') or ''),source_url=list_url,apply_url=link,verification='官方公开职位接口已核验',evidence_file=str(path.relative_to(ROOT)),official_job_id=j['Id'],job_category=j.get('ClassificationOne'),published_at=j.get('PostDate'),notes='用人法人以具体JD/网申为准；完整任职要求见详情。'+('链接直达职位。' if sid=='cicc' else '链接进入官方列表；按职位名或职位ID定位。'))

# SCZQ: all five advertised HTML pages, including historical cohorts.
for j in json.loads((DATA/'sczq-jobs.json').read_text()):
 desc=j.get('description',''); req=desc.split('任职资格：')[-1].split('现在申请')[0].strip() if '任职资格：' in desc else ''
 add(id=key(j['apply_url']),source_id='sczq',org='首创证券招聘体系',org_unit='',title=j['title'],locations=j['locations_text'],recruitment_cycle='2027校招' if '2027' in j['title'] else '2026或更早',requirements=req,source_url=j['list_url'],apply_url=j['apply_url'],verification='官方职位详情已读',evidence_file=j.get('raw_detail',''),notes='包含关联机构，具体用人法人需看岗位详情。')

if (DATA/'pab-jobs.json').exists():
 for j in json.loads((DATA/'pab-jobs.json').read_text()):
  req=j.get('qualification','')
  add(id=key('pab|'+j['idPosition']),source_id='pab',org='平安银行',org_unit=j.get('deptShowName',''),org_level=level(j.get('deptShowName','')),title=j['positionName'],locations=j.get('workCity',''),recruitment_cycle='2027校招' if '2027' in req else '校招（届别待核实）',degree=j.get('education',''),requirements=req,duties=j.get('duty',''),source_url=SOURCES['pab']['url'],apply_url=SOURCES['pab']['url'],verification='官方公开职位接口已核验',evidence_file='data/research/2026-10-13/pab-jobs.json',official_job_id=j['idPosition'],notes='官方校招全职列表；在列表按岗位名称和分行定位，具体毕业窗口见任职要求。')

# Bank campus lists, separated by the source's own unit and stable job identifier.
for sid in ['icbc','bocom','cmbc','cib','cms']:
 path=DATA/(sid+'-jobs.json')
 if not path.exists():continue
 for j in json.loads(path.read_text()):
  org={'icbc':'工商银行','bocom':'交通银行','cmbc':'民生银行','cib':'兴业银行','cms':'招商证券'}[sid]
  req='';duty='';degree='';cycle='校招（届别待核实）';notes='';rawdate=''
  if sid=='icbc':
   ident=j['postId'];title=j.get('publishPostName') or j['postName'];unit=j.get('belongsStruName','');loc=j.get('placeStr','');rawdate=j.get('enterEndTime');duty=strip(j.get('postDepict'))
   notes='官方校园招聘列表；尚未提取完整任职要求。2026年发布不能单独证明毕业届别；多数该批岗位已截止。'
  elif sid=='bocom':
   ident=str(j['positionId']);title=j['pubName'];unit=' / '.join(dict.fromkeys([j.get('bankName',''),j.get('deptName','')]));loc=j.get('workPlace','');rawdate=j.get('endDate')
   cycle=(j.get('projectName') or '校园招聘')+'（毕业届别待核实）';notes='保留官方招聘批次名称；2026年秋季指招录时间，不等于2026届。列表未含完整任职要求。'
  elif sid=='cmbc':
   ident=j['id'];title=j['careerRecruitment_career_name'];unit=j.get('careerRecruitment_career_enterprise_name','');loc=j.get('careerRecruitment_regions_name','');rawdate=j.get('careerRecruitment_career_expirationDate');notes='官方校园招聘列表，包含总行、分行及关联公司；尚未提取完整任职要求。'
  elif sid=='cib':
   ident=j['positionId'];title=j['positionName'];unit=' / '.join(dict.fromkeys(filter(None,[j.get('firstBusinessUnitDesc'),j.get('businessUnitDesc'),j.get('departmentDesc')])));loc=j.get('positionAddr','');rawdate=j.get('expiryDate');req=strip(j.get('positionRequirment'))+' 专业要求：'+strip(j.get('majorRequirment'));duty=strip(j.get('jobDuty'));cycle='2027校招' if '2027' in req else cycle
   notes='官网校园招聘列表；截至日期来自expiryDate，接口期限开关含义未确认，提交前复核。接口分页存在重复，跨轮按官方职位ID去重补齐，覆盖数量见覆盖与缺口。'
  else:
   ident=j['postId'];title=j['postName'];unit=j.get('company','');loc=j.get('workPlaceStr','');rawdate=j.get('endDate');degree=j.get('educationStr','');cycle='2027校招' if '2027' in j.get('projectName','') else '2026或更早' if '2026' in j.get('projectName','') else cycle
   notes='官方校园列表；2030等远期日期待核实，未作为真实投递期限。匿名接口canDelivery标志不可据此判定岗位已关闭；完整条件需打开岗位。'
  d=deadline(rawdate);unit_level=level(unit)
  if sid=='cmbc' and any(x in unit for x in ['民生金租','民生加银','民生科技','民生理财','民生置业']):unit_level='子公司'
  if sid=='cib' and any(x in unit for x in ['兴业国际信托','兴业金融租赁','兴业消费金融','兴银理财']):unit_level='子公司'
  if sid=='cms' and any(x in unit for x in ['招商证券投资有限公司','招商证券国际有限公司','招商致远资本']):unit_level='子公司'
  if sid=='icbc' and any(x in unit for x in ['工银理财','工银金融租赁','工银安盛','工银瑞信','工银科技']):unit_level='子公司'
  add(id=key(sid+'|'+str(ident)),source_id=sid,org=org,org_unit=unit,org_level=unit_level,title=title,locations=loc,recruitment_cycle=cycle,deadline=d,deadline_text=str(rawdate or '未公布') if d else '待核实；接口原值：'+str(rawdate or '未公布'),degree=degree,requirements=req,duties=duty,source_url=SOURCES[sid]['url'],apply_url=SOURCES[sid]['url'],verification='官方公开职位列表已核验',evidence_file=str(path.relative_to(ROOT)),official_job_id=ident,notes=notes)

# Expand rowspans before extraction so each subsidiary/branch retains its own jobs.
sid='research-d5fbc1fb82bc'
for row in table(sid,2)[1:]:
 if len(row)==5 and row[0].isdigit():
  page_record(sid,'浦发银行招聘体系',row[2],row[1],row[4],degree=row[3],dd='2026-10-08',dt='2026年10月8日18:00',apply='https://job.spdb.com.cn',notes='该批公告已截止，保留用于补录/后续批次跟踪；不得视为仍可申请。')
sid='research-55b6099415f9'
for row in table(sid,0)[1:]:
 if len(row)==5:page_record(sid,'兴业证券集团',row[2],row[1],row[3],req=row[4],apply='https://xyzqcampus.hotjob.cn/',notes='类别：'+row[0])
sid='research-26826beb2395'
for row in table(sid,0):
 if not (DATA/'cms-jobs.json').exists() and len(row)==5 and row[0].isdigit():page_record(sid,'招商证券',row[3],' / '.join(dict.fromkeys(row[1:3])),row[4],apply='https://cms.hotjob.cn')
sid='research-5ff13caf8426'
for row in table(sid,0)[1:]:
 title,_,loc=row[1].partition('工作地点：')
 page_record(sid,'南方基金',title.strip(),row[2],loc,req=row[4],degree='硕士及以上' if '硕士' in row[4] else '',dt='10月25日（公告未在该句明示年份）',apply='https://wecruit.hotjob.cn/SU6138665dbef57c3b63841399/pb/school.html')
sid='research-6cc64ad1c2e5'
for row in table(sid,0)[1:]:
 page_record(sid,'常熟农商银行',row[1],'',row[4],req=row[3],dd='2026-10-24',dt='2026年10月24日12:00',apply='https://campus.51job.com/csrcbank/')
sid='research-03954b2b0de6'
for row in table(sid,0)[1:]:page_record(sid,'中银证券',row[0],'',row[1],dd='2026-10-09',dt='2026年10月9日24:00',apply='https://campus.chinahr.com/pages/2027-boc')
sid='research-0d4f42419e78'
for row in table(sid,0)[1:]:page_record(sid,'博时基金',row[2],'',row[6],req=row[7],degree=row[8],kind='岗位类别',apply='https://bosera.hotjob.cn/')
sid='research-acf192465697'
for row in table(sid,0)[1:]:page_record(sid,'中央结算公司',row[1],row[0],row[2],req=row[3],kind='岗位类别')
sid='research-fc1782bbf16e'
for row in table(sid,0)[1:]:page_record(sid,'中国人保财险',''+row[0],'东营市分公司','山东东营',req=row[1],dt='集团整体至2027-04-25；各岗位招满即停，须核实所属机构期限',apply='https://picc.zhiye.com/xiaoyuan')
sid='research-de168ecd7bc6'
for t in [table(sid,i) for i in range(9)]:
 row=t[1] if t[0][0]=='职位名称' else t[0]
 if len(row)>=3:page_record(sid,'江苏金融租赁',row[0],'','南京',degree=row[1],req=' '.join(' '.join(r) for r in t[1:]),apply='https://campus.51job.com/jsleasing/')
sid='research-29c9ad750383'
for index,unit in enumerate(['总公司','北京分公司','上海分公司','银保信科技（北京）有限公司'],1):
 for row in table(sid,index)[1:]:
  if len(row)==6 and row[0].isdigit():
   page_record(sid,'中国银保信',row[2]+'（'+row[1]+'）',unit,row[5],req=row[3],kind='岗位类别' if '/' in row[2] or '软件研发岗 网络' in row[2] else '具体岗位',apply='https://campus.51job.com/cbit2027/job.html')

# Read official static announcements and preserve precise vs partial deadline evidence.
for title,unit,loc,req in [('金融培训生','总行营业部及上海/杭州/绍兴分行','宁波、上海、杭州、绍兴','本科及以上，专业不限'),('科技定向生','总行信息科技部','宁波','本科及以上，计算机、软件工程、人工智能、大数据、算法等相关专业')]:
 add(source_id='ncbank',org='宁波通商银行',org_unit=unit,org_level=level(unit),title=title,locations=loc,recruitment_cycle='2027校招',deadline_text='11月30日（原文未明确年份）',requirements=req,source_url=SOURCES['ncbank']['url'],apply_url=SOURCES['ncbank']['url'],verification='官方公告已读',evidence_file='data/research/2026-10-13/pages/ncbank.html',notes='邮件投递至公告所列邮箱；请先核实完整条件。')
for title in ['国际财富管理及卓越理财部','私人银行','企业及机构银行业务部','资本市场及证券服务部','环球支付方案部','环球贸易方案部']:
 add(source_id='hsbc-2027',org='汇丰银行中国',title='2027管理培训生—'+title,record_type='岗位类别',locations='上海、北京、广州、深圳及中国内地其他城市',recruitment_cycle='2027校招',deadline='2026-10-16',deadline_text='2026年10月16日',source_url=SOURCES['hsbc-2027']['url'],apply_url=SOURCES['hsbc-2027']['url'],verification='招聘专题页已读',evidence_file='data/research/2026-10-13/pages/hsbc-2027.html')
for j in json.loads((ROOT/'examples/boc-2027-jobs.json').read_text()):
 add(source_id='boc',org=j['org'],org_unit=j['org_unit'],title=j['title'],locations='、'.join(j['locations']),recruitment_cycle='2027校招',deadline=j['deadline'],source_url=j['evidence'][0]['url'],apply_url=j['apply_url'],verification='官方公告已读',notes=j['extraction_note'],evidence_file=j['evidence'][0]['raw_file'])

# Aggregate records remain recruitment leads: don't invent job counts by splitting keyword lists.
for j in json.loads((DATA/'radar-finance.json').read_text()):
 add(record_type='聚合招聘线索',org=j['c'],title=j['p'],locations=j['l'],recruitment_cycle=j['w'],deadline_text=j['d'],degree=j.get('e',''),source_id='radar',source_url='https://github.com/jiabaobei/xiaozhao-radar',apply_url=j['u'],verification='聚合记录未回查官方',notes='聚合快照更新时间2026-09-03；招满即止不等于今天仍开放；名称保留原文，可能有简称/错字。',evidence_file='data/research/2026-10-13/radar-jobs.json')

# Research index: keep announcement projects distinct from concrete jobs.
excluded={'research-ef12be2c65ae','research-30196b4db15c','research-16e45b1a622d','research-7389276f78eb','research-c7b1a4910f6d'}
for sid,r in PAGES.items():
 if not r.get('has_2027') or sid in excluded:continue
 text=(ROOT/r['text_file']).read_text()
 if '.edu.cn' not in urlsplit(r['url']).hostname and 'ncss.cn' not in r['url']:continue
 title=r['title']
 if title in ('华中科技大学就业信息网','浙江大学就业服务平台',''):
  lines=[l.strip() for l in text.splitlines() if '2027' in l and ('招聘' in l or '计划' in l) and len(l)<130]
  if not lines:continue
  title=lines[0]
 org=re.split('2027|202 7',title)[0].strip(' -|｜') or title
 add(record_type='招聘公告/项目',source_id=sid,org=org,title=title,recruitment_cycle='2027招聘公告',source_url=r['url'],apply_url=r['url'],verification='高校/就业平台公告已读，需官方复核',evidence_file=r['raw_file'],notes='项目级记录，不计作具体岗位数；机构名为公告标题提取，集团/分支简称可能重复。')

# Merge exact identifiers, retaining multi-source evidence; do not collapse distinct units.
unique={}
for r in records:
 if r['id'] in unique:
  unique[r['id']].setdefault('additional_sources',[]).append(r['source_url'])
 else:unique[r['id']]=r
records=list(unique.values())
records.sort(key=lambda r:({'公告期限内（投递前复核）':0,'截止/在招状态待核实':1,'已截止':2,'历史届别':3}.get(r['status'],4),r['org'],r['title']))
(OUT/'岗位与线索总表.json').write_text(json.dumps(records,ensure_ascii=False,indent=2)+'\n')
fields=['id','record_type','org','org_unit','org_level','track','title','locations','recruitment_cycle','status','deadline','deadline_text','degree','requirements','duties','source_url','apply_url','verification','notes','evidence_file']
labels=['记录ID','记录类型','招聘机构/品牌','招聘单位/部门','机构层级','赛道','职位/项目','地点','届别/批次','状态','截止日期','截止原文','学历','任职要求','工作职责','证据来源','投递/查看入口','核验状态','备注','本地证据']
def safe(value):
 s=str(value or '')
 return "'"+s if s.startswith(('=','+','-','@')) else s
with (OUT/'岗位与线索总表.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.writer(f);w.writerow(labels)
 for r in records:w.writerow([safe(r.get(k,'')) for k in fields])
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
from openpyxl.utils import get_column_letter
wb=Workbook();wb.remove(wb.active)
for name,subset in [('全部记录',records),('具体岗位',[r for r in records if r['record_type']=='具体岗位']),('期限内',[r for r in records if r['status']=='公告期限内（投递前复核）']),('聚合线索',[r for r in records if r['record_type']=='聚合招聘线索'])]:
 ws=wb.create_sheet(name);ws.append(labels)
 for r in subset:ws.append([safe(r.get(k,''))[:32000] for k in fields])
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for cell in ws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='163B55')
 for i in range(1,len(fields)+1):ws.column_dimensions[get_column_letter(i)].width=28 if i not in (7,14,15,16,17,19) else 55
wb.save(OUT/'金融校招岗位大全.xlsx')
with sqlite3.connect(OUT/'catalog.sqlite3') as db:
 db.execute('CREATE TABLE IF NOT EXISTS opportunities(id TEXT PRIMARY KEY, record_type TEXT, org TEXT, status TEXT, body TEXT)')
 db.execute('DELETE FROM opportunities')
 db.executemany('INSERT INTO opportunities VALUES(?,?,?,?,?)',[(r['id'],r['record_type'],r['org'],r['status'],json.dumps(r,ensure_ascii=False)) for r in records])
stats=dict(as_of=TODAY.isoformat(),records=len(records),types=dict(Counter(r['record_type'] for r in records)),statuses=dict(Counter(r['status'] for r in records)),tracks=dict(Counter(r['track'] for r in records)),source_counts=dict(Counter(r['source_id'] for r in records)),exact_labels=len(set(r['org'] for r in records)),deduplication='稳定官方职位ID或机构+单位+岗位+地点+届别；跨源语义重复尚未全部人工合并')
(OUT/'统计.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(stats,ensure_ascii=False,indent=2))
